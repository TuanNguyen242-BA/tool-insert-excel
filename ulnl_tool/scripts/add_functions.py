#!/usr/bin/env python3
"""
add_functions.py - Thêm danh sách chức năng vào sheet 'NL Chuc nang'.

Hỗ trợ:
- Multiple sections (Web, Mobile, Unit test...) mỗi section có row header riêng (cam)
- Nested parent-children đa cấp (đệ quy)
- Excel outline grouping (nút +/-)
- Unit test row tự động tính 30%*60% nỗ lực Phát triển của các mã CN_WEB*/NVJ*/CN_MOBILE Server*
- 3 row tổng (MH, MD, MM)
- 2 row hướng dẫn đánh giá % tái sử dụng

Cách dùng:
  python3 add_functions.py <input_xlsx> <output_xlsx> <functions_json>

Format JSON (mới - sections):
  {
    "sections": [
      {
        "code": "A",
        "name": "NỖ LỰC CHỨC NĂNG WEB",
        "groups": [
          {
            "code": "I", "name": "Quản trị hệ thống",
            "subgroups": [
              {"code": "I.1", "name": "Đăng nhập", "items": [...]}
            ]
          }
        ]
      },
      {
        "code": "B",
        "name": "NỖ LỰC CHỨC NĂNG MOBILE",
        "groups": [...]
      },
      {
        "code": "C",
        "name": "NỖ LỰC UNIT TEST",
        "unit_test": true   # special section - tự sinh 1 row UT1
      }
    ]
  }

Backward-compatible: vẫn nhận format cũ {"groups": [...]} → wrap vào sections[0].

Format item leaf:
  {"name": "X", "ma_cn": "CN_WEB2", "o": 0, "p": 0, "q": 0,
   "mota": "...", "thue_ngoai": {"gp":"Có","pt":"Có","kt":"Có"},
   "ghi_chu": "", "tsd_tu": ""}

Format item cha (recursive):
  {"name": "X", "children": [...]}
"""
import json, sys
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

FONT = 'Times New Roman'
thin = Side(border_style='thin', color='000000')
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

GREEN_LT = 'CCFFCC'
GREEN_DK = '66CC66'
ORANGE_LT = 'FDE9D9'
GRAY_LT = 'F2F2F2'
NAVY = '1F4E79'

LEVEL_SECTION = 0      # Section A/B/C (cam, không thuộc outline)
LEVEL_GROUP = 1
LEVEL_SUBGROUP = 2
LEVEL_PARENT_BASE = 3


def font(size=10, bold=False, italic=False, color='000000'):
    return Font(name=FONT, size=size, bold=bold, italic=italic, color=color)


def fill(hex):
    return PatternFill('solid', start_color=hex, end_color=hex)


def align(h='center', v='center', wrap=True):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)


def write_data_row(ws, r, stt, name, mota, ma_cn, o, p, q, thue_ngoai, ghi_chu, tsd_tu, outline_level):
    """Ghi 1 data row đầy đủ 28 cột công thức."""
    ws.cell(r, 1, stt)
    ws.cell(r, 2, f'=CONCATENATE("CN_",A{r})')
    ws.cell(r, 3, name)
    if mota:
        ws.cell(r, 4, mota)
    else:
        ws.cell(r, 4, f'=CONCATENATE(C{r},CHAR(10),"Dự kiến: ...")')
    ws.cell(r, 5, f'=SUM(F{r}:H{r})')
    ws.cell(r, 6, f'=K{r}*(1-$O{r})')
    ws.cell(r, 7, f'=L{r}*(1-$P{r})')
    ws.cell(r, 8, f'=M{r}*(1-$Q{r})')
    ws.cell(r, 9, ma_cn)
    ws.cell(r, 10, f"=IF(I{r}=\"\",\"\",VLOOKUP(I{r},'Định nghĩa Loại Chức năng'!$B$4:$C$92,2,FALSE))")
    ws.cell(r, 11, f"=IF(I{r}=\"\",\"\",VLOOKUP(I{r}&$K$7,CSDL_CN!$D$4:$E$264,2,FALSE))")
    ws.cell(r, 12, f"=IF(I{r}=\"\",\"\",VLOOKUP(I{r}&$L$7,CSDL_CN!$D$4:$E$264,2,FALSE))")
    ws.cell(r, 13, f"=IF(I{r}=\"\",\"\",VLOOKUP(I{r}&$M$7,CSDL_CN!$D$4:$E$264,2,FALSE))")
    ws.cell(r, 14, f'=SUM(K{r}:M{r},0)')
    ws.cell(r, 15, o)
    ws.cell(r, 16, p)
    ws.cell(r, 17, q)
    ws.cell(r, 18, tsd_tu or '')
    ws.cell(r, 19, ghi_chu or '')
    tn = thue_ngoai or {}
    ws.cell(r, 20, tn.get('gp', 'Có'))
    ws.cell(r, 21, tn.get('pt', 'Có'))
    ws.cell(r, 22, tn.get('kt', 'Có'))
    ws.cell(r, 23, f'=IF(T{r}="Có",K{r}*(1-$O{r})*0.37,"")')
    ws.cell(r, 24, f'=IF(U{r}="Có",L{r}*(1-$P{r}),"")')
    ws.cell(r, 25, f'=IF(V{r}="Có",M{r}*(1-$Q{r}),"")')
    ws.cell(r, 26, f'=IF(K{r}="","",IF(W{r}="",K{r}*(1-$O{r}),K{r}*(1-$O{r})-W{r}))')
    ws.cell(r, 27, f'=IF(L{r}="","",IF(X{r}="",L{r}*(1-$P{r}),L{r}*(1-$P{r})-X{r}))')
    ws.cell(r, 28, f'=IF(M{r}="","",IF(Y{r}="",M{r}*(1-$Q{r}),M{r}*(1-$Q{r})-Y{r}))')

    for ci in range(1, 29):
        c = ws.cell(r, ci)
        c.border = BORDER
        c.font = font(10)
        h = 'center' if ci in (1, 2, 9, 20, 21, 22) else 'left' if ci in (3, 4, 18, 19) else 'right'
        c.alignment = align(h, 'top')
        if ci in (5, 6, 7, 8, 11, 12, 13, 14, 23, 24, 25, 26, 27, 28):
            c.number_format = '0.00'
        if ci in (15, 16, 17):
            c.number_format = '0%'
    ws.cell(r, 9).font = font(10, bold=True, color=NAVY)
    for ci in (15, 16, 17):
        ws.cell(r, ci).fill = fill('FFFFFF')

    if outline_level > 0:
        ws.row_dimensions[r].outline_level = outline_level


def write_parent_row(ws, r, stt, name, outline_level):
    """Tính năng cha (cam nhạt)."""
    ws.cell(r, 1, stt)
    ws.cell(r, 2, f'=CONCATENATE("CN_",A{r})')
    ws.cell(r, 3, name)
    for ci in range(1, 29):
        c = ws.cell(r, ci)
        c.fill = fill(ORANGE_LT)
        c.border = BORDER
        c.font = font(10, bold=True)
        h = 'center' if ci in (1, 2) else 'left'
        c.alignment = align(h, 'center')
    if outline_level > 0:
        ws.row_dimensions[r].outline_level = outline_level


def write_group_row(ws, r, code, name, outline_level):
    """Group / Subgroup row (xanh nhạt)."""
    ws.cell(r, 1, code)
    ws.cell(r, 1).alignment = align('center', 'center')
    if name:
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        ws.cell(r, 2, name)
        ws.cell(r, 2).alignment = align('left', 'center')
    for ci in range(1, 29):
        c = ws.cell(r, ci)
        c.fill = fill(GREEN_LT)
        c.border = BORDER
        c.font = font(11, bold=True)
    if outline_level > 0:
        ws.row_dimensions[r].outline_level = outline_level


def write_section_row(ws, r, code, name):
    """Section row (cam đậm) - như row 8 'A NỖ LỰC CHỨC NĂNG'."""
    ws.cell(r, 1, code)
    ws.cell(r, 1).font = font(12, bold=True)
    ws.cell(r, 1).fill = fill(ORANGE_LT)
    ws.cell(r, 1).alignment = align('center', 'center')
    ws.cell(r, 1).border = BORDER
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    ws.cell(r, 2, name)
    ws.cell(r, 2).font = font(12, bold=True)
    ws.cell(r, 2).fill = fill(ORANGE_LT)
    ws.cell(r, 2).alignment = align('center', 'center')
    ws.cell(r, 2).border = BORDER
    for ci in range(4, 29):
        c = ws.cell(r, ci)
        c.fill = fill(ORANGE_LT)
        c.border = BORDER
    ws.row_dimensions[r].height = 22


def write_ut_row(ws, r, data_start_row):
    """Row 'Thực hiện Unit test' - mã UNITTEST, công thức 30%*60% PT của Web+Mobile Server."""
    ws.cell(r, 1, 514)
    ws.cell(r, 2, 'UT1')
    ws.cell(r, 3, 'Thực hiện Unit test')
    ws.cell(r, 4, ('Nỗ lực thực hiện Unit Test:\n'
                   '- Đảm bảo code coverage tối thiểu 50% đối với code mới\n'
                   '- Chỉ phát sinh nỗ lực của đầu việc phát triển (sau tái sử dụng)'))
    ws.cell(r, 5, f'=SUM(F{r}:H{r})')
    # G (Phát triển sau TSD) - SUMIFS công thức
    g_formula = (
        f'=(SUMIFS($G${data_start_row}:$G${r-1},$I${data_start_row}:$I${r-1},"=CN_WEB*")'
        f'+SUMIFS($G${data_start_row}:$G${r-1},$I${data_start_row}:$I${r-1},"=NVJ*")'
        f'+SUMIFS($G${data_start_row}:$G${r-1},$I${data_start_row}:$I${r-1},"=CN_MOBILE Server*"))*0.3*0.6'
    )
    ws.cell(r, 7, g_formula)
    # F, H trống (không phát sinh GP/KT cho UT)
    ws.cell(r, 9, 'UNITTEST')
    ws.cell(r, 10, f"=IF(I{r}=\"\",\"\",VLOOKUP(I{r},'Định nghĩa Loại Chức năng'!$B$4:$C$92,2,FALSE))")
    # L (Phát triển chưa TSD)
    l_formula = (
        f'=(SUMIFS($L${data_start_row}:$L${r-1},$I${data_start_row}:$I${r-1},"=CN_WEB*")'
        f'+SUMIFS($L${data_start_row}:$L${r-1},$I${data_start_row}:$I${r-1},"=NVJ*")'
        f'+SUMIFS($L${data_start_row}:$L${r-1},$I${data_start_row}:$I${r-1},"=CN_MOBILE Server*"))*0.3*0.6'
    )
    ws.cell(r, 12, l_formula)
    # N = SUM K:M (chỉ có L)
    ws.cell(r, 14, f'=SUM(K{r}:M{r},0)')
    # O, P, Q = 0 (UT không tái sử dụng)
    ws.cell(r, 15, 0)
    ws.cell(r, 16, 0)
    ws.cell(r, 17, 0)
    # T, U, V: chỉ U có ý nghĩa (Phát triển → Unit test → có thuê ngoài)
    ws.cell(r, 20, '')
    ws.cell(r, 21, 'Có')
    ws.cell(r, 22, '')
    # W, X, Y
    ws.cell(r, 23, f'=IF(T{r}="Có",F{r}*1,"")')
    ws.cell(r, 24, f'=IF(U{r}="Có",G{r},"")')
    ws.cell(r, 25, f'=IF(V{r}="Có",H{r}*1,"")')

    for ci in range(1, 29):
        c = ws.cell(r, ci)
        c.border = BORDER
        c.font = font(10)
        h = 'center' if ci in (1, 2, 9, 20, 21, 22) else 'left' if ci in (3, 4, 18, 19) else 'right'
        c.alignment = align(h, 'top')
        if ci in (5, 6, 7, 8, 11, 12, 13, 14, 23, 24, 25):
            c.number_format = '0.00'
        if ci in (15, 16, 17):
            c.number_format = '0%'
    ws.cell(r, 9).font = font(10, bold=True, color=NAVY)


def write_total_mh_row(ws, r, start_row):
    """'Nỗ lực (MH)' - xanh đậm."""
    ws.cell(r, 2, 'Nỗ lực (MH)')
    for col, letter in [(5,'E'),(6,'F'),(7,'G'),(8,'H'),(11,'K'),(12,'L'),
                        (13,'M'),(14,'N'),(23,'W'),(24,'X'),(25,'Y'),
                        (26,'Z'),(27,'AA'),(28,'AB')]:
        ws.cell(r, col, f'=SUM({letter}{start_row}:{letter}{r-1})')
        ws.cell(r, col).number_format = '0.00'
    for ci in range(1, 29):
        c = ws.cell(r, ci)
        c.font = font(11, bold=True)
        c.fill = fill(GREEN_DK)
        c.border = BORDER
        c.alignment = align('center', 'center')


def write_md_row(ws, r):
    """'Nỗ lực (MD)' = MH/8 - xám."""
    ws.cell(r, 2, 'Nỗ lực (MD)')
    for col, letter in [(5,'E'),(6,'F'),(7,'G'),(8,'H'),(11,'K'),(12,'L'),
                        (13,'M'),(14,'N'),(23,'W'),(24,'X'),(25,'Y'),
                        (26,'Z'),(27,'AA'),(28,'AB')]:
        ws.cell(r, col, f'={letter}{r-1}/8')
        ws.cell(r, col).number_format = '0.00'
    for ci in range(1, 29):
        c = ws.cell(r, ci)
        c.font = font(11, bold=True)
        c.fill = fill(GRAY_LT)
        c.border = BORDER
        c.alignment = align('center', 'center')


def write_mm_row(ws, r):
    """'Nỗ lực (MM)' = MD/22 - xám."""
    ws.cell(r, 2, 'Nỗ lực (MM)')
    ws.cell(r, 5, f'=E{r-1}/22')
    ws.cell(r, 5).number_format = '0.00'
    for ci in range(1, 29):
        c = ws.cell(r, ci)
        c.font = font(11, bold=True)
        c.fill = fill(GRAY_LT)
        c.border = BORDER
        c.alignment = align('center', 'center')


def write_huong_dan(ws, r):
    """Row hướng dẫn header + body."""
    ws.cell(r, 2, 'Hướng dẫn đánh giá % tái sử dụng:')
    for ci in range(1, 29):
        c = ws.cell(r, ci)
        c.font = font(11, bold=True, italic=True)
        c.alignment = align('left', 'center')

    ws.cell(r+1, 2, ('Quản trị dự án đánh giá xem:\n'
                     '* Có chức năng tương tự trong hệ thống không?\n'
                     '* Có chức năng tương tự trong hệ thống khác của đơn vị không?\n'
                     '* Có tái sử dụng được thiết kế của các chức năng tương tự không?'))
    for ci in range(1, 29):
        c = ws.cell(r+1, ci)
        c.font = font(10, italic=True)
        c.alignment = align('left', 'top', wrap=True)
    ws.row_dimensions[r+1].height = 60


def write_items(ws, items, current_row, stt_path, outline_level):
    """Ghi đệ quy danh sách items."""
    for idx, item in enumerate(items, 1):
        stt = f'{stt_path}.{idx}' if stt_path else str(idx)
        has_children = 'children' in item and item['children']
        
        if has_children:
            write_parent_row(ws, current_row, stt, item['name'], outline_level)
            current_row += 1
            current_row = write_items(ws, item['children'], current_row, stt, outline_level + 1)
        else:
            if 'ma_cn' not in item:
                print(f"WARNING: '{item.get('name', '?')}' thiếu ma_cn, skip")
                continue
            write_data_row(ws, current_row, stt,
                           item['name'], item.get('mota'), item['ma_cn'],
                           item.get('o', 0), item.get('p', 0), item.get('q', 0),
                           item.get('thue_ngoai'),
                           item.get('ghi_chu'), item.get('tsd_tu'),
                           outline_level)
            current_row += 1
    return current_row


def add_functions(input_xlsx, output_xlsx, functions_json):
    with open(functions_json, 'r', encoding='utf-8') as f:
        spec = json.load(f)

    wb = load_workbook(input_xlsx)
    ws = wb['NL Chuc nang']

    # Backward compatibility
    if 'sections' not in spec and 'groups' in spec:
        spec = {'sections': [{'code': 'A', 'name': 'NỖ LỰC CHỨC NĂNG', 'groups': spec['groups']}]}

    # Xóa demo rows từ row 8 trở đi (clear cả row 8 vì giờ section header được render bởi section[0])
    if ws.max_row >= 8:
        for r in range(ws.max_row, 7, -1):
            ws.delete_rows(r)

    cur_row = 8
    data_start_row = 9  # row đầu tiên sau section A header

    # Write sections
    ut_row_idx = None
    for sec_idx, section in enumerate(spec['sections']):
        # Section row (cam) - giống row 8 cũ
        write_section_row(ws, cur_row, section['code'], section['name'])
        cur_row += 1

        # Special: Unit test section
        if section.get('unit_test'):
            ut_row_idx = cur_row
            write_ut_row(ws, cur_row, data_start_row)
            cur_row += 1
            continue

        # Standard section
        for grp in section.get('groups', []):
            write_group_row(ws, cur_row, grp['code'], grp['name'], LEVEL_GROUP)
            cur_row += 1
            if 'subgroups' in grp and grp['subgroups']:
                for sub in grp['subgroups']:
                    write_group_row(ws, cur_row, sub['code'], sub['name'], LEVEL_SUBGROUP)
                    cur_row += 1
                    cur_row = write_items(ws, sub.get('items', []), cur_row, '', LEVEL_PARENT_BASE)
            elif 'items' in grp:
                cur_row = write_items(ws, grp['items'], cur_row, '', LEVEL_SUBGROUP)

    # Add total + hướng dẫn
    total_mh_row = cur_row
    write_total_mh_row(ws, total_mh_row, data_start_row)
    cur_row += 1
    write_md_row(ws, cur_row)
    cur_row += 1
    write_mm_row(ws, cur_row)
    cur_row += 1
    # Spacer
    cur_row += 1
    write_huong_dan(ws, cur_row)
    last_row = cur_row + 1

    ws.sheet_properties.outlinePr.summaryBelow = False
    ws.sheet_properties.outlinePr.summaryRight = False

    print(f"DATA_START_ROW={data_start_row}")
    print(f"TOTAL_MH_ROW={total_mh_row}")
    if ut_row_idx:
        print(f"UT_ROW={ut_row_idx}")
    print(f"LAST_ROW={last_row}")

    wb.save(output_xlsx)
    return total_mh_row, ut_row_idx


if __name__ == '__main__':
    if len(sys.argv) < 4:
        print(__doc__)
        sys.exit(1)
    add_functions(sys.argv[1], sys.argv[2], sys.argv[3])
