from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


RAW_TABLE_LIST_HEADERS = [
    "STT",
    "Loại bảng",
    "Tên bảng đích",
    "Mô tả bảng đích",
    "Tần suất\n(batch/streaming)",
    "Đường dẫn trên kho dữ liệu",
    "Các trường Partition",
    "Thời gian lưu trữ",
    "Định dạng lưu trữ",
    "Ghi chú",
]

RAW_MAPPING_HEADERS = [
    "STT",
    "Tên hệ thống nguồn",
    "Loại bảng",
    "Tên bảng đích",
    "Đường dẫn lưu trữ vùng dữ liệu thô (Raw Zone)",
    "Tên bảng nguồn",
    "Tên topic Kafka",
    "Mô tả yêu cầu tổng hợp bảng trên Work Zone (các bước, điều kiện)",
    "Ghi chú",
]

WORK_TABLE_LIST_HEADERS = [
    "STT",
    "Loại bảng",
    "Tên bảng đích",
    "Mô tả bảng đích",
    "Tần suất\n(batch/streaming)",
    "Lĩnh vực cấp 1",
    "Lĩnh vực cấp 2",
    "Lĩnh vực cấp 3",
    "Đường dẫn trên kho dữ liệu",
    "Các trường Partition",
    "Thời gian lưu trữ",
    "Định dạng lưu trữ",
    "Ghi chú",
]

WORK_MAPPING_HEADERS = [
    "STT",
    "Tên hệ thống nguồn",
    "Loại bảng",
    "Tên bảng đích",
    "Đường dẫn lưu trữ đạt chất lượng vùng dữ liệu làm việc (Work Zone)",
    "Đường dẫn lưu trữ chưa đạt chất lượng vùng dữ liệu làm việc (Work Zone - Bad data)",
    "Tên bảng nguồn",
    "Đường dẫn lưu trữ bảng nguồn tại vùng dữ liệu thô (Raw Zone)",
    "Mô tả các bước thực hiện tổng hợp bảng",
    "Ghi chú",
]

RELATION_HEADERS = [
    "STT",
    "Tên bảng đích",
    "Bảng chính",
    "Trường chính",
    "Bảng liên kết",
    "Trường liên kết (FK)",
]

FLOW_HEADERS = [
    "STT",
    "Tên bảng đích",
    "Trường đích",
    "Định dạng",
    "Tên bảng đích (chuẩn hóa)",
    "Trường đích (chuẩn hóa)",
    "Định dạng (chuẩn hóa)",
    "Ý nghĩa trường đích",
    "Hệ thống nguồn",
    "Tên bảng nguồn",
    "Trường nguồn",
    "Mô tả trường nguồn",
    "Nullable",
    "Unique",
    "Key",
    "Mặc định",
    "Loại bảng",
    "Mô tả bảng",
    "Mã kiểm tra",
    "Mã chuẩn hóa",
    "Phương pháp kiểm tra chi tiết",
    "Phương pháp chuẩn hóa chi tiết",
    "Giá trị mặc định",
    "Hệ thống nguồn",
    "Schema.Table",
    "Source Field Name",
    "ETL Rules",
]

LOG_HEADERS = [
    "Ngày thay đổi",
    "Vị trí",
    "A/M/D",
    "Nguồn gốc (lý do) tạo mới/sửa đổi/xóa",
    "Đầu mối yêu cầu",
    "Mô tả thay đổi",
    "Người thực hiện",
    "Phiên bản",
]


DEFAULT_OPTIONS = {
    "domain": "<phụ nữ>",
    "source_system_default": "Hệ thống kho dữ liệu của TCCT",
    "raw_zone": "raw_zone",
    "work_zone": "work_zone",
    "good_data_folder": "good_data",
    "bad_data_path": "work_zone/{domain}/bad_data/bad_data_catalog",
    "raw_topic_prefix": "dbz.{domain}",
    "retention": "Vĩnh viễn",
    "storage_format": "Parquet",
    "batch_frequency": "Xử lý toàn bộ dữ liệu một lần duy nhất dữ liệu ngày N",
    "stream_frequency": "Xử lý dữ liệu liên tục (streaming) trong ngày và chạy 1 lần/ngày",
    "field_domain_l1": "Lĩnh vực Quân chúng",
    "master_prefix": "master_",
    "event_prefix": "event_",
    "reference_prefix": "ref_",
    "add_lake_updated_time": True,
    "lake_updated_field": "lake_updated_time",
    "lake_updated_format": "timestamp",
    "lake_updated_description": "Hệ thống tự sinh - Thời gian cập nhật",
}


@dataclass
class SourceField:
    order: Any = ""
    table_name: str = ""
    field_name: str = ""
    description: str = ""
    datatype: str = ""
    unique: str = ""
    nullable: str = ""
    key: str = ""
    default: str = ""
    source_system: str = ""
    source_table: str = ""
    source_field: str = ""
    ref_table: str = ""
    ref_field: str = ""
    etl_rules: str = ""
    dqc_code: str = ""
    std_code: str = ""
    check_detail: str = ""
    std_detail: str = ""
    note: str = ""
    table_type: str = ""
    table_desc: str = ""


@dataclass
class SourceTable:
    order: Any
    name: str
    table_type: str
    description: str
    source_system: str
    fields: list[SourceField]


def normalized_text(value: Any) -> str:
    text = str(value or "").strip().lower()
    text = unicodedata.normalize("NFD", text)
    text = "".join(ch for ch in text if unicodedata.category(ch) != "Mn")
    text = text.replace("đ", "d")
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


HEADER_ALIASES = {
    "stt": "order",
    "bang dich": "table_name",
    "ten bang dich": "table_name",
    "truong dich": "field_name",
    "y nghia": "description",
    "y nghia truong dich": "description",
    "dinh dang": "datatype",
    "unique": "unique",
    "nullable": "nullable",
    "pk fk": "key",
    "key": "key",
    "mac dinh": "default",
    "he thong nguon": "source_system",
    "bang nguon": "source_table",
    "ten bang nguon": "source_table",
    "truong nguon": "source_field",
    "source field name": "source_field",
    "mo ta truong nguon": "description",
    "bang tham chieu": "ref_table",
    "truong tham chieu": "ref_field",
    "etl rules": "etl_rules",
    "ma kiem tra": "dqc_code",
    "ma chuan hoa": "std_code",
    "phuong phap kiem tra chi tiet": "check_detail",
    "phuong phap chuan hoa chi tiet": "std_detail",
    "ghi chu": "note",
    "loai bang": "table_type",
    "mo ta bang": "table_desc",
    "mo ta bang dich": "table_desc",
}


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def list_sheet_names(path: Path) -> list[str]:
    wb = load_workbook(path, read_only=True, data_only=True)
    names = list(wb.sheetnames)
    wb.close()
    return names


def _find_header_row(ws) -> tuple[int, dict[int, str]]:
    for row_idx in range(1, min(ws.max_row, 30) + 1):
        mapping = {}
        normalized_headers = []
        for col_idx in range(1, min(ws.max_column, 80) + 1):
            key = HEADER_ALIASES.get(normalized_text(ws.cell(row_idx, col_idx).value))
            if key:
                mapping[col_idx] = key
                normalized_headers.append(key)
        if {"table_name", "field_name"}.issubset(set(normalized_headers)):
            return row_idx, mapping
    raise ValueError("Khong tim thay dong header co cot 'Bang dich' va 'Truong dich'.")


def parse_source_workbook(path: Path, sheet_name: str) -> list[SourceTable]:
    wb = load_workbook(path, read_only=False, data_only=True)
    if sheet_name not in wb.sheetnames:
        wb.close()
        raise ValueError(f"Sheet khong ton tai: {sheet_name}")
    ws = wb[sheet_name]
    header_row, header_map = _find_header_row(ws)

    fields: list[SourceField] = []
    carry: dict[str, str] = {
        "order": "",
        "table_name": "",
        "table_type": "",
        "table_desc": "",
        "source_system": "",
        "source_table": "",
    }
    for row_idx in range(header_row + 1, ws.max_row + 1):
        values: dict[str, str] = {}
        for col_idx, attr in header_map.items():
            values[attr] = _cell_text(ws.cell(row_idx, col_idx).value)

        if not any(values.values()):
            continue
        for attr in carry:
            if values.get(attr):
                carry[attr] = values[attr]
            elif attr in {"order", "table_name", "table_type", "table_desc", "source_system"}:
                values[attr] = carry[attr]

        if not values.get("table_name") or not values.get("field_name"):
            continue

        source_table = values.get("source_table") or values.get("table_name") or carry.get("source_table", "")
        source_field = values.get("source_field") or values.get("field_name")
        fields.append(SourceField(
            order=values.get("order", ""),
            table_name=values.get("table_name", ""),
            field_name=values.get("field_name", ""),
            description=values.get("description", ""),
            datatype=values.get("datatype", ""),
            unique=values.get("unique", ""),
            nullable=values.get("nullable", ""),
            key=values.get("key", ""),
            default=values.get("default", ""),
            source_system=values.get("source_system", ""),
            source_table=source_table,
            source_field=source_field,
            ref_table=values.get("ref_table", ""),
            ref_field=values.get("ref_field", ""),
            etl_rules=values.get("etl_rules", ""),
            dqc_code=values.get("dqc_code", ""),
            std_code=values.get("std_code", ""),
            check_detail=values.get("check_detail", ""),
            std_detail=values.get("std_detail", ""),
            note=values.get("note", ""),
            table_type=values.get("table_type", ""),
            table_desc=values.get("table_desc", ""),
        ))

    wb.close()
    table_map: dict[str, SourceTable] = {}
    for fld in fields:
        if fld.table_name not in table_map:
            table_map[fld.table_name] = SourceTable(
                order=fld.order,
                name=fld.table_name,
                table_type=fld.table_type,
                description=fld.table_desc or fld.description,
                source_system=fld.source_system,
                fields=[],
            )
        table = table_map[fld.table_name]
        if not table.table_type and fld.table_type:
            table.table_type = fld.table_type
        if not table.description and fld.table_desc:
            table.description = fld.table_desc
        if not table.source_system and fld.source_system:
            table.source_system = fld.source_system
        table.fields.append(fld)
    return list(table_map.values())


def coerce_options(raw: dict | None = None) -> dict:
    options = dict(DEFAULT_OPTIONS)
    for key, value in (raw or {}).items():
        if key not in options:
            continue
        if isinstance(options[key], bool):
            options[key] = str(value).lower() in {"1", "true", "yes", "on"}
        else:
            options[key] = str(value)
    return options


def table_kind(table_type: str) -> str:
    text = normalized_text(table_type)
    if "su kien" in text or "event" in text:
        return "event"
    if "tham chieu" in text or "reference" in text:
        return "reference"
    return "master"


def work_table_name(table: SourceTable, options: dict) -> str:
    prefix = {
        "master": options["master_prefix"],
        "event": options["event_prefix"],
        "reference": options["reference_prefix"],
    }[table_kind(table.table_type)]
    return f"{options['domain']}_{prefix}{table.name}"


def frequency_for(table: SourceTable, options: dict) -> str:
    if table_kind(table.table_type) == "event":
        return options["stream_frequency"]
    return options["batch_frequency"]


def raw_path(table: SourceTable, options: dict) -> str:
    return f"{options['raw_zone']}/{options['domain']}/{table.name}"


def work_good_path(table: SourceTable, options: dict) -> str:
    return f"{options['work_zone']}/{options['domain']}/{options['good_data_folder']}/{work_table_name(table, options)}"


def work_bad_path(options: dict) -> str:
    return options["bad_data_path"].format(domain=options["domain"])


def source_system(table: SourceTable, options: dict) -> str:
    return table.source_system or options["source_system_default"]


def kafka_topic(table: SourceTable, options: dict) -> str:
    prefix = options["raw_topic_prefix"].format(domain=options["domain"])
    return f"{prefix}.{table.name}"


def _base_workbook() -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Log"
    _write_log_sheet(ws)
    return wb


def _styles():
    thin = Side(style="thin", color="000000")
    return {
        "header_fill": PatternFill("solid", fgColor="FFFF00"),
        "cyan_fill": PatternFill("solid", fgColor="00FFFF"),
        "orange_fill": PatternFill("solid", fgColor="FF9900"),
        "gray_fill": PatternFill("solid", fgColor="C0C0C0"),
        "system_fill": PatternFill("solid", fgColor="EDEDED"),
        "header_font": Font(name="Times New Roman", bold=True, size=11),
        "body_font": Font(name="Times New Roman", size=11),
        "center": Alignment(horizontal="center", vertical="center", wrap_text=True),
        "left": Alignment(horizontal="left", vertical="top", wrap_text=True),
        "border": Border(left=thin, right=thin, top=thin, bottom=thin),
    }


def _write_log_sheet(ws) -> None:
    st = _styles()
    ws.cell(1, 1, "A - Tạo mới, M - Sửa đổi, D - Xóa bỏ")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(LOG_HEADERS))
    for col, header in enumerate(LOG_HEADERS, 1):
        cell = ws.cell(2, col, header)
        cell.font = st["header_font"]
        cell.alignment = st["center"]
        cell.border = st["border"]
    for row in range(3, 23):
        for col in range(1, len(LOG_HEADERS) + 1):
            ws.cell(row, col).border = st["border"]
    widths = [13, 32, 8, 36, 18, 42, 18, 42]
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A3"


def _write_table(ws, headers: list[str], rows: list[list[Any]], colored: dict[int, str] | None = None) -> None:
    st = _styles()
    colored = colored or {}
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(1, col_idx, header)
        cell.font = st["header_font"]
        cell.alignment = st["center"]
        cell.fill = {
            "cyan": st["cyan_fill"],
            "orange": st["orange_fill"],
            "gray": st["gray_fill"],
        }.get(colored.get(col_idx), st["header_fill"])
        cell.border = st["border"]
    for row_idx, row in enumerate(rows, 2):
        is_system = bool(row and str(row[0]).startswith("__system__"))
        if is_system:
            row = row[1:]
        for col_idx, value in enumerate(row, 1):
            cell = ws.cell(row_idx, col_idx, value)
            cell.font = st["body_font"]
            cell.alignment = st["left"]
            cell.border = st["border"]
            if is_system:
                cell.fill = st["system_fill"]
    max_row = max(1, len(rows) + 1)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max_row}"
    _fit_columns(ws, headers)


def _fit_columns(ws, headers: list[str]) -> None:
    widths = {
        "STT": 6,
        "Loại bảng": 26,
        "Tên bảng đích": 34,
        "Trường đích": 28,
        "Định dạng": 18,
        "Ý nghĩa trường đích": 42,
        "Mô tả trường nguồn": 42,
        "Mô tả bảng": 34,
        "Mô tả bảng đích": 34,
        "Đường dẫn trên kho dữ liệu": 44,
        "Tên hệ thống nguồn": 28,
        "Tên topic Kafka": 34,
        "Phương pháp kiểm tra chi tiết": 52,
        "Phương pháp chuẩn hóa chi tiết": 52,
        "Ghi chú": 28,
        "Ten topic Kafka": 34,
        "Phuong phap kiem tra chi tiet": 52,
        "Phuong phap chuan hoa chi tiet": 52,
        "Ghi chu": 28,
    }
    for idx, header in enumerate(headers, 1):
        single_line = str(header).split("\n", 1)[0]
        ws.column_dimensions[get_column_letter(idx)].width = widths.get(single_line, min(max(len(single_line) + 4, 14), 38))


def _relation_rows(tables: list[SourceTable], options: dict, use_work_table: bool) -> list[list[Any]]:
    rows = []
    for idx, table in enumerate(tables, 1):
        name = work_table_name(table, options) if use_work_table else table.name
        rows.append([idx, name, table.name, "", "", ""])
    return rows


def _raw_flow_rows(tables: list[SourceTable], options: dict) -> list[list[Any]]:
    rows = []
    table_no = 0
    for table in tables:
        table_no += 1
        for fld in table.fields:
            rows.append([
                table_no,
                table.name,
                fld.field_name,
                fld.datatype,
                "",
                "",
                "",
                fld.description,
                source_system(table, options),
                fld.source_table or table.name,
                fld.source_field or fld.field_name,
                fld.description,
                fld.nullable,
                fld.unique,
                fld.key,
                fld.default,
                table.table_type,
                table.description,
                "",
                source_system(table, options),
                "",
                fld.source_field or fld.field_name,
                fld.etl_rules,
            ])
    return rows


def _work_flow_rows(tables: list[SourceTable], options: dict) -> list[list[Any]]:
    rows = []
    table_no = 0
    for table in tables:
        table_no += 1
        for fld in table.fields:
            rows.append([
                table_no,
                table.name,
                fld.field_name,
                fld.datatype,
                "",
                "",
                "",
                fld.description,
                source_system(table, options),
                fld.source_table or table.name,
                fld.source_field or fld.field_name,
                fld.description,
                fld.nullable,
                fld.unique,
                fld.key,
                fld.default,
                table.table_type,
                table.description,
                fld.dqc_code,
                fld.std_code,
                fld.check_detail,
                fld.std_detail,
                "",
                source_system(table, options),
                "",
                fld.source_field or fld.field_name,
                fld.etl_rules,
            ])
        if options["add_lake_updated_time"]:
            rows.append([
                "__system__",
                table_no,
                table.name,
                options["lake_updated_field"],
                options["lake_updated_format"],
                "",
                "",
                "",
                options["lake_updated_description"],
                source_system(table, options),
                "N/A",
                "N/A",
                "N/A",
                "NO",
                "",
                "",
                "",
                table.table_type,
                table.description,
                "",
                "",
                "",
                "",
                "",
                source_system(table, options),
                "",
                "",
                "",
            ])
    return rows


def build_raw_workbook(tables: list[SourceTable], output_path: Path, options: dict | None = None) -> Path:
    options = coerce_options(options)
    wb = _base_workbook()

    table_rows = []
    for idx, table in enumerate(tables, 1):
        table_rows.append([
            idx,
            table.table_type,
            table.name,
            table.description,
            frequency_for(table, options),
            raw_path(table, options),
            "",
            options["retention"],
            options["storage_format"],
            "",
        ])
    _write_table(wb.create_sheet("Table List"), RAW_TABLE_LIST_HEADERS, table_rows)

    mapping_rows = []
    for idx, table in enumerate(tables, 1):
        mapping_rows.append([
            idx,
            source_system(table, options),
            table.table_type,
            table.name,
            raw_path(table, options),
            table.name,
            kafka_topic(table, options),
            f"Thực hiện tích hợp dữ liệu của bảng {table.name}",
            "",
        ])
    _write_table(wb.create_sheet("Mapping_Source_Raw"), RAW_MAPPING_HEADERS, mapping_rows)
    _write_table(wb.create_sheet("Bảng thông tin liên kết dữ liệu"), RELATION_HEADERS, _relation_rows(tables, options, False))
    _write_table(
        wb.create_sheet("Bảng tổng hợp phân tích luồng"),
        FLOW_HEADERS,
        _raw_flow_rows(tables, options),
        colored={5: "cyan", 6: "cyan", 7: "cyan", 21: "orange", 22: "orange", 23: "gray", 24: "gray", 25: "gray", 26: "gray", 27: "gray"},
    )
    wb.save(output_path)
    return output_path


def build_work_workbook(tables: list[SourceTable], output_path: Path, options: dict | None = None) -> Path:
    options = coerce_options(options)
    wb = _base_workbook()

    table_rows = []
    for idx, table in enumerate(tables, 1):
        table_rows.append([
            idx,
            table.table_type,
            work_table_name(table, options),
            table.description,
            frequency_for(table, options),
            options["field_domain_l1"],
            "",
            "",
            work_good_path(table, options),
            "",
            options["retention"],
            options["storage_format"],
            "",
        ])
    _write_table(wb.create_sheet("Table List"), WORK_TABLE_LIST_HEADERS, table_rows)

    mapping_rows = []
    detail = (
        "Thực hiện mapping 1-1 theo từng trường từ Raw Zone lên Work Zone\n\n"
        "B1: Dữ liệu đi qua các quy tắc kiểm tra (DQC) cho một số trường được chỉ định\n"
        "+ Nếu không đạt -> trả toàn bộ bản ghi về Bad Zone\n"
        "+ Nếu đạt -> tiến hành đến bước 2\n\n"
        "B2: Dữ liệu đi qua các quy tắc chuẩn hóa (STD) cho một số trường được chỉ định\n"
        "+ Nếu chuẩn hóa thành công -> Tiếp tục các bước xử lý tiếp theo\n"
        "+ Nếu chuẩn hóa thất bại -> trả toàn bộ bản ghi về Bad Zone\n\n"
        "B3: Tiến hành biến đổi tên trường tên bảng theo như thông tin chuẩn hóa và lưu xuống HDFS\n\n"
        "Tham khảo chi tiết ở sheet [Bảng tổng hợp phân tích luồng] rule xử lý cho từng trường trong từng bảng"
    )
    for idx, table in enumerate(tables, 1):
        mapping_rows.append([
            idx,
            source_system(table, options),
            table.table_type,
            work_table_name(table, options),
            work_good_path(table, options),
            work_bad_path(options),
            table.name,
            raw_path(table, options),
            detail,
            "Tiến hành biến đổi tên bảng theo nhu cầu chuẩn hóa",
        ])
    _write_table(wb.create_sheet("Mapping_Raw_Work"), WORK_MAPPING_HEADERS, mapping_rows)
    _write_table(wb.create_sheet("Bảng thông tin liên kết dữ liệu"), RELATION_HEADERS, _relation_rows(tables, options, True))
    _write_table(
        wb.create_sheet("Bảng tổng hợp phân tích luồng"),
        FLOW_HEADERS,
        _work_flow_rows(tables, options),
        colored={5: "cyan", 6: "cyan", 7: "cyan", 21: "orange", 22: "orange", 23: "gray", 24: "gray", 25: "gray", 26: "gray", 27: "gray"},
    )
    wb.save(output_path)
    return output_path


def preview_tables(tables: list[SourceTable], options: dict | None = None, limit: int = 20) -> dict:
    options = coerce_options(options)
    rows = []
    field_count = 0
    for idx, table in enumerate(tables, 1):
        field_count += len(table.fields)
        if len(rows) < limit:
            rows.append({
                "stt": idx,
                "table_name": table.name,
                "work_table_name": work_table_name(table, options),
                "table_type": table.table_type,
                "fields": len(table.fields),
                "description": table.description,
                "raw_path": raw_path(table, options),
                "work_path": work_good_path(table, options),
            })
    return {
        "tables": rows,
        "table_count": len(tables),
        "field_count": field_count,
        "truncated": len(tables) > limit,
    }
